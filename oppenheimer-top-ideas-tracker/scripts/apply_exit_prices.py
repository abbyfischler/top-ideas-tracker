#!/usr/bin/env python3
"""
apply_exit_prices.py — Take the Bloomberg-filled spreadsheet
(exit_prices_needed.xlsx, with the "Closing Price" column filled in) and
turn it into data/exit_prices.csv, which build_data.py uses to compute a
real return for the month an analyst switches tickers.

Without this file, a ticker switch shows a blank return for that one month
(there's no way to know what the abandoned position was worth on the day
it was dropped). With it, that month's return reflects the abandoned
position's actual performance up to the switch — which is what feeds into
each analyst's all-time and yearly totals, so filling this in changes those
numbers too (up or down, depending on whether the exited position had
gained or lost ground).

USAGE
    python scripts/apply_exit_prices.py exit_prices_needed.xlsx
    python scripts/apply_exit_prices.py exit_prices_needed.xlsx --allow-partial

Then rebuild as usual:
    python scripts/build_data.py
    python scripts/build_site.py
"""
import argparse
import csv
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('xlsx_path')
    ap.add_argument('--allow-partial', action='store_true',
                     help="Don't stop if some rows are still missing a price — just use what's filled in.")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.xlsx_path, data_only=True)
    ws = wb['Prices to Look Up'] if 'Prices to Look Up' in wb.sheetnames else wb.active

    rows = []
    missing = []
    for r in range(2, ws.max_row + 1):
        ticker = ws.cell(row=r, column=1).value
        date_needed = ws.cell(row=r, column=3).value
        price = ws.cell(row=r, column=4).value
        if not ticker:
            continue
        if price is None or price == '':
            missing.append((ticker, date_needed))
            continue
        try:
            price = float(price)
        except (TypeError, ValueError):
            missing.append((ticker, date_needed))
            continue
        # normalize the date to YYYY-MM-DD to match the tracker's internal format
        if hasattr(date_needed, 'strftime'):
            date_str = date_needed.strftime('%Y-%m-%d')
        else:
            # handle a plain typed string like "5/18/2026"
            from datetime import datetime
            date_str = datetime.strptime(str(date_needed), '%m/%d/%Y').strftime('%Y-%m-%d')
        rows.append((ticker, date_str, price))

    print(f"Filled in: {len(rows)} rows")
    if missing:
        print(f"Still blank: {len(missing)} rows")
        for t, d in missing[:15]:
            print(f"   - {t} on {d}")
        if len(missing) > 15:
            print(f"   ...and {len(missing)-15} more")
        if not args.allow_partial:
            print("\nNot writing exit_prices.csv yet — fill in the remaining rows, or re-run with "
                  "--allow-partial to use just what's filled in (blank ones stay unfixed for now).")
            sys.exit(1)

    out_path = ROOT / "data" / "exit_prices.csv"
    with open(out_path, 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(['ticker', 'date', 'price'])
        for t, d, p in rows:
            w.writerow([t, d, p])

    print(f"\nWrote {out_path} ({len(rows)} prices)")
    print("Now run: python scripts/build_data.py && python scripts/build_site.py")


if __name__ == "__main__":
    main()
